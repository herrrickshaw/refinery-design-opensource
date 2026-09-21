"""Fetch ExxonMobil's public crude-oil assay workbooks and distil them into
``refinery_design/data/assays.json``.

Only property *values* are extracted (whole-crude properties, per-cut yields
and qualities, the TBP curve) - the workbooks themselves are not
redistributed.  ExxonMobil publishes them "courtesy of ExxonMobil" with an
explicit no-warranty notice; see ``docs/VENDOR_REFERENCE.md``.

Run: python scripts/build_assay_data.py   (needs: pip install openpyxl)
"""
from __future__ import annotations

import io
import json
import re
import urllib.request
from pathlib import Path

import openpyxl

BASE = "https://corporate.exxonmobil.com/-/media/global/files/crude-oils/xls/"
# name -> published path fragment (some assays live under a year folder)
CRUDES = {
    "bakken": "bakken.xlsx",
    "azeri_btc": "2024/azeri_btc.xlsx",
    "qua_iboe": "2024/qua_iboe.xlsx",
    "upper_zakum": "upper_zakum.xlsx",
    "alaska_north_slope": "alaska_north_slope.xlsx",
    "dalia": "dalia.xlsx",
    "cold_lake_blend": "2024/cold_lake_blend.xlsx",
    "kearl": "2024/kearl.xlsx",
}

ROWS = {  # workbook row label -> key
    "Yield (% wt)": "wt_pct", "Yield (% vol)": "vol_pct",
    "Density @ 15°C (g/cc)": "density", "API Gravity": "api", "UOPK": "uop_k",
    "Molecular Weight (g/mol)": "mw", "Total Sulfur (% wt)": "sulfur_wt",
    "Basic Nitrogen (ppm)": "basic_n_ppm", "Total Nitrogen (ppm)": "total_n_ppm",
    "Total Acid Number (mgKOH/g)": "tan", "Pour Point (°C)": "pour_C",
    "Hydrogen (% wt)": "hydrogen_wt", "Micro Carbon Residue (% wt)": "mcr_wt",
    "Vanadium (ppm)": "v_ppm", "Nickel (ppm)": "ni_ppm",
    "C7 Asphaltenes (% wt)": "c7_asphaltenes_wt", "Volume Average B.P. (°C)": "vabp_C",
}


def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def parse(raw: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    rows = {}
    name = origin = None
    for row in ws.iter_rows():
        cells = {c.column_letter: c.value for c in row if c.value is not None}
        if not cells:
            continue
        first = next(iter(cells.values()))
        label = str(first).strip()
        if label == "Crude:":
            name = str(list(cells.values())[1]).strip()
        if label == "Origin:":
            origin = str(list(cells.values())[1]).strip()
        rows[label] = cells
    start, end = rows["Start (°C)"], rows["End (°C) "[:-1]]
    cols = [c for c in start if c not in ("B",)]  # C = whole crude, D.. = cuts
    def col_of(label):
        return rows[label]
    whole = {}
    for label, key in ROWS.items():
        if label in rows and "C" in rows[label]:
            whole[key] = _num(rows[label]["C"])
    cuts = []
    for c in cols[1:]:
        if not isinstance(start[c], (int, float, str)) or str(start[c]).strip() == "":
            continue
        cut = {"start_C": start[c], "end_C": end[c]}
        for label, key in ROWS.items():
            if label in rows and c in rows[label]:
                cut[key] = _num(rows[label][c])
        cuts.append(cut)
    # TBP curve: sheet 2, column B = temperature (degC), C = cumulative wt%,
    # D = cumulative vol% distilled (the clean table starts below the
    # 0/10/20... interval grid).
    tbp = []
    for row in wb.worksheets[1].iter_rows(min_row=70):
        b, c, d = (row[i].value for i in (1, 2, 3))
        if all(isinstance(v, (int, float)) for v in (b, c, d)):
            tbp.append([float(b), round(float(c), 3), round(float(d), 3)])
    return {"name": name, "origin": origin, "whole": whole, "cuts": cuts, "tbp_T_cumwt_cumvol": tbp}


def main() -> None:
    out = {"_source": "ExxonMobil crude oil assays (public downloads), corporate.exxonmobil.com/"
                      "what-we-do/energy-supply/crude-trading/crude-oil-assays; values extracted, "
                      "no warranty (see docs/VENDOR_REFERENCE.md)", "crudes": {}}
    for key, frag in CRUDES.items():
        req = urllib.request.Request(BASE + frag, headers={"User-Agent": "Mozilla/5.0"})
        out["crudes"][key] = parse(urllib.request.urlopen(req, timeout=60).read())
        print("parsed", key)
    dest = Path(__file__).resolve().parents[1] / "refinery_design" / "data" / "assays.json"
    dest.write_text(json.dumps(out, indent=1, ensure_ascii=False))
    print("wrote", dest)


if __name__ == "__main__":
    main()
